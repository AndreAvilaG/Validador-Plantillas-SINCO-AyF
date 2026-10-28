"""Pruebas de la información bancaria en Terceros (Etapa 1).

Regla (pedida por Andrea): Banco, Tipo Cuenta, Número cuenta y Email Tesorería van JUNTAS. Si se diligencia cualquiera de las cuatro, las otras
tres también deben estar diligenciadas; si no hay información bancaria, las cuatro quedan vacías.

Es UN solo error por fila: su columna nombra las que faltan unidas con « / » y su mensaje las lista todas. Así el reporte (PDF y pantalla) junta en
un solo punto las filas con el mismo faltante (antes salía una tarjeta por columna: la misma lista de 98 filas se repetía tres veces) y el Excel
marcado pinta cada celda vacía.

Qué se protege:
  * La regla funciona con las DOS versiones de nombres de la plantilla: Banco / Tipo Cuenta / Número cuenta / Email Tesorería y
    TerBanco / TerTipoCuentaB / TerCuentaB / TerEmailTes.
  * Lo que ya funcionaba no cambia: una fila con las cuatro llenas o las cuatro vacías no da ningún error nuevo, y el código de banco inexistente
    se sigue avisando.
  * La pantalla muestra el aviso y el Excel marcado deja marcada cada celda vacía.

Cómo correrlas (desde la raíz del repositorio):  pip install pytest   y luego   pytest pruebas
Los datos son inventados: no se deben agregar archivos de clientes a este repositorio.
"""
import io
import os
import re

import pytest

from _cargador import cargar, ejecutar, llenar_empresa

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CLAVE = "Información bancaria incompleta."

# (banco, tipo de cuenta, número de cuenta, correo de tesorería) en cada versión de la plantilla
VERSION_B = ("Banco", "Tipo Cuenta", "Número cuenta", "Email Tesorería")
VERSION_A = ("TerBanco", "TerTipoCuentaB", "TerCuentaB", "TerEmailTes")
VERSIONES = [pytest.param(VERSION_B, id="Banco/Tipo Cuenta/Número cuenta/Email Tesorería"),
             pytest.param(VERSION_A, id="TerBanco/TerTipoCuentaB/TerCuentaB/TerEmailTes")]
VALORES = ("51", "2", "138100051844", "tesoreria@empresa.com")           # un valor válido para cada una de las cuatro columnas


def _validar(filas, columnas):
    """Corre validar_terceros del validador real sobre filas [{columna: valor}]; devuelve solo los errores de la información bancaria."""
    import pandas as pd
    df = pd.DataFrame([{"Tipo identificación": "NIT", "Número de identificación": str(900100200 + i), **{c: None for c in columnas}, **f}
                       for i, f in enumerate(filas)])

    def correr(ns):
        tipos, bancos, ciudades, ciiu = ns["cargar_catalogos"]()
        return ns["validar_terceros"](df, tipos, bancos, ciudades, ciiu)

    errores = ejecutar(["validar_terceros", "cargar_catalogos"], correr)
    return [e for e in errores if e[4] == CLAVE]


def _fila(columnas, llenas):
    """Una fila con llenas las columnas cuyo índice (0 a 3) se indica y vacías las demás."""
    return {c: VALORES[i] for i, c in enumerate(columnas) if i in llenas}


def _faltantes(errores):
    """Las columnas que faltan según los errores (cada error trae las suyas unidas con « / »)."""
    return sorted(c for e in errores for c in e[1].split(" / "))


# ───────────────────────── Lo que debe pasar sin aviso ─────────────────────────
@pytest.mark.parametrize("columnas", VERSIONES)
def test_las_cuatro_vacias_no_dan_aviso(columnas):
    assert _validar([_fila(columnas, set())], columnas) == []


@pytest.mark.parametrize("columnas", VERSIONES)
def test_las_cuatro_llenas_no_dan_aviso(columnas):
    assert _validar([_fila(columnas, {0, 1, 2, 3})], columnas) == []


# ───────────────────────── Lo que debe avisar: cualquier combinación incompleta ─────────────────────────
@pytest.mark.parametrize("columnas", VERSIONES)
@pytest.mark.parametrize("llenas", [{0}, {1}, {2}, {3}, {0, 1}, {0, 3}, {1, 2}, {2, 3}, {0, 1, 2}, {0, 1, 3}, {0, 2, 3}, {1, 2, 3}])
def test_cualquier_combinacion_incompleta_avisa_exactamente_las_columnas_que_faltan(columnas, llenas):
    errores = _validar([_fila(columnas, llenas)], columnas)
    assert _faltantes(errores) == sorted(c for i, c in enumerate(columnas) if i not in llenas)
    assert len(errores) == 1                                                 # un solo error por fila, aunque falten varias columnas


@pytest.mark.parametrize("columnas", VERSIONES)
def test_solo_el_correo_diligenciado_pide_banco_tipo_y_numero(columnas):
    """El caso reportado: el archivo traía todos los correos pero ninguna otra columna bancaria."""
    errores = _validar([_fila(columnas, {3})], columnas)
    assert _faltantes(errores) == sorted(columnas[:3]) and len(errores) == 1
    assert errores[0][1] == " / ".join(columnas[:3])


@pytest.mark.parametrize("columnas", VERSIONES)
def test_banco_tipo_y_numero_sin_correo_pide_el_correo(columnas):
    errores = _validar([_fila(columnas, {0, 1, 2})], columnas)
    assert [e[1] for e in errores] == [columnas[3]]


@pytest.mark.parametrize("columnas", VERSIONES)
def test_el_error_es_bloqueante_en_la_fila_correcta_y_el_mensaje_explica_que_falta(columnas):
    filas = [_fila(columnas, {0, 1, 2, 3}), _fila(columnas, {0}), _fila(columnas, set())]
    errores = _validar(filas, columnas)
    assert {e[0] for e in errores} == {3}                                    # fila 3 de Excel = segunda fila de datos (la 1 es el encabezado)
    assert len(errores) == 1                                                 # la fila 3 falla una sola vez, no una por cada columna que falta
    unico = errores[0]
    assert unico[1] == " / ".join(columnas[1:]) and unico[2] == "formato" and unico[3].startswith("❌")
    assert f"tiene diligenciado {columnas[0]} pero faltan {columnas[1]}, {columnas[2]} y {columnas[3]}." in unico[3]
    assert f"({columnas[0]}, {columnas[1]}, {columnas[2]} y {columnas[3]}), debes diligenciar todas" in unico[3]


@pytest.mark.parametrize("columnas", VERSIONES)
def test_cada_fila_se_evalua_por_separado(columnas):
    errores = _validar([_fila(columnas, {0}), _fila(columnas, {0, 1, 2, 3}), _fila(columnas, {3}), _fila(columnas, set())], columnas)
    por_fila = {}
    for e in errores:
        por_fila.setdefault(e[0], []).append(e[1])
    assert sorted(por_fila) == [2, 4] and len(por_fila[2]) == 1 and len(por_fila[4]) == 1
    assert por_fila[2] == [" / ".join(columnas[1:])] and por_fila[4] == [" / ".join(columnas[:3])]


@pytest.mark.parametrize("vacio", ["", "   ", "nan", "NaN", "None", None])
@pytest.mark.parametrize("columnas", VERSIONES)
def test_un_valor_en_blanco_cuenta_como_vacio(columnas, vacio):
    """Un espacio o un «nan» de Excel no es información: si las otras tres están llenas, falta la cuarta."""
    fila = {**_fila(columnas, {0, 1, 2}), columnas[3]: vacio}
    assert [e[1] for e in _validar([fila], columnas)] == [columnas[3]]


@pytest.mark.parametrize("columnas", VERSIONES)
def test_un_valor_con_espacios_alrededor_cuenta_como_lleno(columnas):
    fila = {c: f"  {v}  " for c, v in zip(columnas, VALORES)}
    assert _validar([fila], columnas) == []


# ───────────────────────── Nombres de columna ─────────────────────────
def test_los_nombres_se_reconocen_sin_importar_mayusculas_ni_espacios_de_mas():
    columnas = (" BANCO ", "tipo cuenta", "NÚMERO CUENTA", "email tesorería")
    assert _faltantes(_validar([_fila(columnas, {0})], columnas)) == sorted(c for c in columnas[1:])


def test_tambien_se_acepta_el_nombre_sin_tilde_del_correo_y_del_numero():
    columnas = ("Banco", "Tipo Cuenta", "Numero cuenta", "Email Tesoreria")
    assert _faltantes(_validar([_fila(columnas, {0})], columnas)) == sorted(columnas[1:])


def test_si_el_archivo_no_trae_alguna_de_las_columnas_el_grupo_son_las_que_si_trae():
    """Solo se validan las columnas presentes (como el resto de la validación): sin la columna del correo, el grupo son las otras tres."""
    tres = VERSION_B[:3]
    assert _faltantes(_validar([_fila(tres, {0})], tres)) == sorted(tres[1:])
    assert _validar([_fila(tres, {0, 1, 2})], tres) == []


def test_si_solo_hay_una_columna_del_grupo_no_hay_con_que_compararla():
    assert _validar([{"Banco": "51"}], ("Banco",)) == []
    assert _validar([{"Banco": "51"}], ()) == []


# ───────────────────────── Lo que ya funcionaba sigue igual ─────────────────────────
@pytest.mark.parametrize("columnas", VERSIONES)
def test_un_banco_inexistente_sigue_avisando_ademas_de_lo_demas(columnas):
    import pandas as pd
    fila = {c: v for c, v in zip(columnas, ("99", "2", "123456789", "a@b.co"))}
    df = pd.DataFrame([{"Tipo identificación": "NIT", "Número de identificación": "900100200", **fila}])

    def correr(ns):
        tipos, bancos, ciudades, ciiu = ns["cargar_catalogos"]()
        return ns["validar_terceros"](df, tipos, bancos, ciudades, ciiu)

    claves = [e[4] for e in ejecutar(["validar_terceros", "cargar_catalogos"], correr)]
    assert CLAVE not in claves                                               # el grupo está completo: sin aviso de información incompleta
    if columnas == VERSION_B:                                                # la regla del código de banco usa el nombre «Banco»
        assert "Banco no válido." in claves


def test_el_aviso_tiene_su_texto_amigable_para_la_pantalla_y_el_pdf():
    ns, *_ = cargar(["FRASES_AMIGABLES"])
    frase = ns["FRASES_AMIGABLES"][CLAVE]
    assert frase["titulo"] == "Información bancaria incompleta" and "Email Tesorería" in frase["ocurre"] and frase["accion"]


# ───────────────────────── El Excel marcado ─────────────────────────
@pytest.mark.parametrize("columnas", VERSIONES)
def test_el_excel_marcado_deja_marcada_cada_celda_vacia(columnas):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIGRACION"
    ws.append(["Tipo identificación", "Número de identificación", *columnas])
    ws.append(["NIT", "900100200", *VALORES])                                # completa: no se marca
    ws.append(["NIT", "900100201", VALORES[0], None, None, None])            # solo el banco: faltan las otras tres
    datos = io.BytesIO()
    wb.save(datos)
    errores = _validar([{c: v for c, v in zip(columnas, VALORES)}, {columnas[0]: VALORES[0]}], columnas)

    def marcar(ns):
        return ns["marcar_errores_excel"](io.BytesIO(datos.getvalue()), errores, None)

    marcado = ejecutar(["marcar_errores_excel"], marcar)
    hoja = openpyxl.load_workbook(io.BytesIO(marcado))["MIGRACION"]
    def marcada(fila, col):
        celda = hoja.cell(row=fila, column=col)
        return bool(celda.comment) or (celda.fill and celda.fill.fill_type == "solid")
    assert [marcada(3, c) for c in (4, 5, 6)] == [True, True, True]          # las tres celdas vacías de la fila 3
    assert not marcada(3, 3) and not any(marcada(2, c) for c in range(3, 7))  # el banco lleno y la fila completa no se marcan
    resumen = list(openpyxl.load_workbook(io.BytesIO(marcado))["Inconsistencias"].iter_rows(values_only=True))
    assert len(resumen) == 2 and resumen[1][0] == 3 and resumen[1][1] == " / ".join(columnas[1:])   # encabezado + UNA fila para la fila 3


# ───────────────────────── El reporte PDF: un solo punto por faltante ─────────────────────────
def _texto_pdf(errores):
    pypdf = pytest.importorskip("pypdf")
    reemplazo = {"generar_pdf_inconsistencias": [("except Exception:\n        return None", "except Exception:\n        raise")]}
    pdf = ejecutar(["generar_pdf_inconsistencias"], lambda ns: ns["generar_pdf_inconsistencias"](
        "Reporte de Inconsistencias — Plantillas Maestras", [("Terceros", errores)], "Cliente de prueba", "999000001"), reemplazos=reemplazo)
    assert pdf and pdf[:4] == b"%PDF"
    return re.sub(r"\s+", " ", " ".join((p.extract_text() or "") for p in pypdf.PdfReader(io.BytesIO(pdf)).pages))


@pytest.mark.parametrize("columnas", VERSIONES)
def test_las_mismas_filas_salen_en_un_solo_punto_del_pdf_y_no_repetidas_por_columna(columnas):
    """Regresión: 98 terceros con solo el correo diligenciado aparecían en TRES tarjetas (Banco, Tipo Cuenta y Número cuenta), cada una con
    «Se repite en 98 filas»: las mismas filas contadas tres veces."""
    errores = _validar([{columnas[3]: f"t{i}@empresa.com"} for i in range(98)], columnas)
    assert len(errores) == 98                                                # antes eran 294 (98 filas × 3 columnas)
    texto = _texto_pdf(errores)
    assert texto.count("Se repite en 98 filas") == 1
    assert texto.count("Información bancaria incompleta") == 1
    assert f"El tercero tiene diligenciado {columnas[3]} pero faltan {columnas[0]}, {columnas[1]} y {columnas[2]}." in texto


@pytest.mark.parametrize("columnas", VERSIONES)
def test_faltantes_distintos_son_puntos_distintos_y_cada_uno_cuenta_sus_filas(columnas):
    filas = [{columnas[3]: "a@b.co"}] * 5 + [{columnas[0]: "51"}] * 3
    texto = _texto_pdf(_validar(filas, columnas))
    assert texto.count("Se repite en 5 filas") == 1 and texto.count("Se repite en 3 filas") == 1
    assert texto.count("Información bancaria incompleta") == 2


# ───────────────────────── En pantalla ─────────────────────────
@pytest.mark.parametrize("columnas", VERSIONES)
def test_la_pantalla_muestra_el_aviso_cuando_hay_informacion_bancaria_incompleta(columnas):
    import openpyxl
    from streamlit.testing.v1 import AppTest
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIGRACION"
    ws.append(["Tipo identificación", "Número de identificación", *columnas])
    ws.append(["NIT", "900100200", *VALORES])
    for i in range(3):                                                       # tres terceros con el mismo faltante: una sola tarjeta
        ws.append(["NIT", f"90010020{i + 1}", None, None, None, "solo.correo@empresa.com"])
    datos = io.BytesIO()
    wb.save(datos)
    app = AppTest.from_file(os.path.join(RAIZ, "validador_sinco_ayf.py"), default_timeout=120)
    app.run()
    next(b for b in app.button if b.label == "➡️ Iniciar validación").click().run()
    radio = app.radio[0]
    radio.set_value([o for o in radio.options if "Ya tengo mis plantillas" in o][0]).run()
    llenar_empresa(app)
    next(f for f in app.file_uploader if f.key == "terceros_e1" or f.key.startswith("terceros_e1_")).set_value([("t.xlsx", datos.getvalue(), XLSX)]).run()
    assert not app.exception, [str(e.value) for e in app.exception]
    texto = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", " ".join(m.value for m in app.markdown)))
    assert texto.count("Información bancaria incompleta") == 1 and "Banco, Tipo Cuenta, Número cuenta y Email Tesorería" in texto
