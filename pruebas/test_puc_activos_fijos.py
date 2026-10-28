"""Pruebas de la validación del grupo 15 (Propiedad, planta y equipo) en el maestro de cuentas contables — PUC Principal y PUC Fiscal.

Regla: en el maestro las cuentas van a un nivel general; el detalle individual de cada bien (su unidad, su ubicación, sus características) se
administra en el submódulo de Activos Fijos. Cuando el cliente crea cuentas con ese detalle, se muestra una RECOMENDACIÓN (advertencia, no bloquea).

Qué se protege:
  * Se marcan las cuentas del grupo 15 cuyo nombre identifica una unidad («LOCAL 4», «APTO 701»), un inmueble por su dirección («CALLE 52 10 20»),
    un bien dentro de un inmueble con nombre («LOCAL 4 EDIFICIO X») o un tipo de bien separado en una cuenta padre que no lo distingue.
  * NO se marcan las cuentas generales ni las que solo contienen una palabra como «oficina» o «garaje» sin ese contexto (EQUIPO DE OFICINA,
    MUEBLES ENSERES OFICINAS, OFICINAS bajo la cuenta OFICINAS, PARQUEADEROS, GARAJES Y DEPOSITOS), ni nada fuera del grupo 15.
  * La observación sale igual en el Excel marcado (celda resaltada con su comentario y fila en la hoja Inconsistencias), en el PDF y en la pantalla;
    los datos originales del cliente no se modifican y las demás validaciones no cambian.

Cómo correrlas (desde la raíz del repositorio):  pip install pytest pypdf   y luego   pytest pruebas
Los datos son inventados: no se deben agregar archivos de clientes a este repositorio.
"""
import io
import os
import re

import pytest

from _cargador import ejecutar

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CLAVE = "Detalle individual del activo en el maestro (grupo 15)"
TITULO = "Cuentas de Propiedad, planta y equipo con detalle individual del activo"
TEXTO_SUGERIDO = ("Se identifica un nivel de detalle específico del activo en el maestro de cuentas contables. Se recomienda validar si esta clasificación "
                  "debe manejarse de forma general en el maestro contable, dejando el detalle individual del bien para su administración en el submódulo "
                  "de Activos Fijos.")

MAESTROS = [("1", "ACTIVO"), ("15", "PROPIEDAD PLANTA Y EQUIPO"), ("1516", "CONSTRUCCIONES Y EDIFICACIONES"), ("151605", "EDIFICIOS"),
            ("151610", "OFICINAS"), ("151670", "PARQUEADEROS, GARAJES Y DEPOSITOS"), ("1524", "EQUIPO DE OFICINA"), ("152405", "MUEBLES Y ENSERES"),
            ("1550", "PROPIEDAD PLANTA Y EQUIPO"), ("155005", "TERRENOS"), ("1555", "PROPIEDAD PLANTA Y EQUIPO DE INVERSION"),
            ("155516", "CONSTRUCCIONES Y EDIFICACIONES"), ("1592", "DEPRECIACION ACUMULADA"), ("159205", "CONSTRUCCIONES Y EDIFICACIONES"),
            ("5", "GASTOS"), ("51", "OPERACIONALES DE ADMINISTRACION"), ("5120", "ARRENDAMIENTOS"), ("512010", "CONSTRUCCIONES Y EDIFICACIONES")]

# Auxiliares que SÍ se deben marcar: (código, descripción)
CON_DETALLE = [
    ("15160505", "APARTAMENTOS"),                        # tipo de bien separado bajo EDIFICIOS
    ("15160510", "LOCAL"),                               # ídem
    ("15500505", "CALLE 52 10 20"),                      # terreno identificado por su dirección
    ("15551610", "APTO 701 TORRE SOL"),                  # unidad concreta
    ("15551620", "LOCAL 4 EDIFICIO CENTRAL"),            # unidad concreta en un edificio con nombre
    ("15551630", "GARAJE 12 CONJUNTO LOS ROBLES"),
    ("15551640", "LOCALES PUNTO 90"),                    # grupo de locales de un proyecto identificado
    ("15551650", "MEJORAS EDIFICIO CRA 7 12 30"),        # inmueble identificado por su dirección
    ("15920505", "DEPRECIACION LOCAL 4"),                # también en la depreciación del bien individual
]
# Auxiliares generales que NO se deben marcar
GENERALES = [
    ("15161005", "OFICINAS"),                            # repite el nombre de su cuenta padre: es el nivel general de esa subcuenta
    ("15167005", "GARAJES Y DEPOSITOS"),                 # el padre ya habla de garajes
    ("15167010", "GARAJES Y DEPSITOS"),                  # con la errata de digitación
    ("15240505", "MUEBLES ENSERES OFICINAS"),            # «oficinas» dentro de una cuenta general de muebles
    ("15240510", "EQUIPO DE OFICINA 2019"),              # un año no es una unidad
    ("15500510", "TERRENOS URBANOS"),
    ("15551605", "CONSTRUCCIONES Y EDIFICACIONES"),
    ("15920510", "DEPRECIACION EDIFICIOS"),
    ("51201005", "LOCAL 2 ARRIENDO"),                    # fuera del grupo 15 la regla no aplica
    ("51201010", "CALLE 10 BODEGA"),
]


def _fila_puc(cols, codigo, desc, tipo):
    return {**{c: "" for c in cols}, "Código": codigo, "Descripción": desc, "Tipo": tipo}


def _validar(auxiliares, es_fiscal=False, devolver_df=False):
    """Corre validar_puc del validador real sobre un PUC mínimo con esos auxiliares. Devuelve las observaciones del grupo 15 (o, con
    devolver_df, también el DataFrame antes y después para comprobar que no se modificó)."""
    import pandas as pd

    def correr(ns):
        cols = list(ns["_COLS_PUC"]) + (["Equivalencia"] if es_fiscal and "Equivalencia" not in ns["_COLS_PUC"] else [])
        df = pd.DataFrame([_fila_puc(cols, c, d, "M") for c, d in MAESTROS] + [_fila_puc(cols, c, d, "A") for c, d in auxiliares])
        antes = df.copy(deep=True)
        errores = ns["validar_puc"](df, es_fiscal=es_fiscal)
        return errores, antes, df

    errores, antes, despues = ejecutar(["validar_puc", "_COLS_PUC"], correr)
    marcadas = [e for e in errores if e[4] == CLAVE]
    return (marcadas, errores, antes, despues) if devolver_df else marcadas


def _codigos(observaciones):
    return sorted(re.search(r'La cuenta "(\d+)"', e[3]).group(1) for e in observaciones)


# ───────────────────────── La lógica de contexto ─────────────────────────
def _motivo(desc, tipo="A", padre=""):
    return ejecutar(["_motivo_detalle_activo"], lambda ns: ns["_motivo_detalle_activo"](desc, tipo, padre))


@pytest.mark.parametrize("desc", ["LOCAL 2 EDIFICIO DAVINCI", "GARAJE 15", "Apto 401 Miramar", "APARTAMENTO 3B", "OFICINA 502", "LOCAL No. 7", "Local Nº 7", "LOCAL #7", "apto. 401",
                                  "LOCALES PUNTO 170", "BODEGA 12", "PARQUEADERO 33", "LOTE 8", "CASA 5", "Consultorio 204", "LOCAL NRO 3"])
def test_un_nombre_que_identifica_una_unidad_concreta_se_marca(desc):
    assert "unidad específica" in _motivo(desc)


@pytest.mark.parametrize("desc", ["CALLE 128", "CLL 95 14 29", "CARRERA 7 12 30", "Avenida 68 Nº 12", "KM 12", "MEJORAS EDIFICIO CLL 95 14 29", "DIAGONAL 40 A",
                                  "Cra. 7 # 12-30", "Calle No. 52", "Calle 10"])
def test_un_nombre_que_identifica_un_inmueble_por_su_direccion_se_marca(desc):
    assert "por su dirección" in _motivo(desc)


@pytest.mark.parametrize("desc", ["LOCAL EDIFICIO DAVINCI", "GARAJES CONJUNTO LOS ROBLES", "OFICINAS CENTRO COMERCIAL ANDINO", "LOCAL TORRE NORTE"])
def test_un_bien_ubicado_en_un_inmueble_con_nombre_se_marca(desc):
    assert "inmueble específico" in _motivo(desc)


@pytest.mark.parametrize("desc,padre", [("APARTAMENTOS", "EDIFICIOS"), ("LOCAL", "EDIFICIOS"), ("OFICINA", "REVALUACIONES ACTIVOS FIJOS"),
                                        ("LOCALES Y OFICINAS", "EDIFICIOS"), ("Bodegas", "CONSTRUCCIONES Y EDIFICACIONES")])
def test_una_auxiliar_que_solo_nombra_un_tipo_de_bien_bajo_un_padre_que_no_lo_distingue_se_marca(desc, padre):
    assert "tipo de activo" in _motivo(desc, "A", padre)


@pytest.mark.parametrize("desc,padre", [("OFICINAS", "OFICINAS"), ("OFICINA", "OFICINAS"), ("GARAJES", "PARQUEADEROS, GARAJES Y DEPOSITOS"),
                                        ("LOCALES", "LOCALES COMERCIALES"), ("APARTAMENTOS", "Apartamentos")])
def test_si_la_cuenta_padre_ya_habla_de_ese_tipo_es_el_nivel_general_y_no_se_marca(desc, padre):
    assert _motivo(desc, "A", padre) is None


@pytest.mark.parametrize("desc", ["EQUIPO DE OFICINA", "MUEBLES ENSERES OFICINAS", "MUEBLES Y ENSERES", "EDIFICIOS", "CONSTRUCCIONES Y EDIFICACIONES",
                                  "GARAJES Y DEPSITOS", "TERRENOS URBANOS", "MAQUINARIA Y EQUIPO CONSTRUCCION",
                                  "AJTE INFLACION DE GARAJES Y DEPOSITOS", "DEPRECIACION EDIFICIOS", "EQUIPO DE OFICINA 2019", "TV 55 PULGADAS", ""])
def test_una_cuenta_general_o_con_solo_una_palabra_parecida_no_se_marca(desc):
    assert _motivo(desc, "A", "EDIFICIOS") is None


def test_un_nombre_con_solo_el_tipo_pero_en_una_cuenta_maestra_no_se_marca():
    """Las subcuentas del PUC estándar (OFICINAS, PARQUEADEROS, GARAJES Y DEPOSITOS…) son maestras: solo las AUXILIARES que parten un tipo de bien
    se consideran detalle."""
    assert _motivo("LOCALES", "M", "EDIFICIOS") is None and _motivo("OFICINAS", "M", "") is None
    assert _motivo("PARQUEADEROS, GARAJES Y DEPOSITOS", "M", "CONSTRUCCIONES Y EDIFICACIONES") is None
    assert _motivo("APARTAMENTOS", "A", "") is None                          # sin cuenta padre conocida no se puede afirmar nada


@pytest.mark.parametrize("desc", ["OFICINA 1999", "OFICINA 2019", "OFICINA 2024", "OFICINA 2025 A"])
def test_un_numero_que_parece_un_ano_no_es_una_unidad(desc):
    assert _motivo(desc) is None


def test_un_numero_que_no_parece_un_ano_si_es_una_unidad():
    assert "unidad específica" in _motivo("OFICINA 1500") and "unidad específica" in _motivo("OFICINA 502")


# ───────────────────────── En la validación del maestro (principal y fiscal) ─────────────────────────
@pytest.mark.parametrize("es_fiscal", [False, True], ids=["PUC Principal", "PUC Fiscal"])
def test_se_marcan_exactamente_las_cuentas_con_detalle_individual(es_fiscal):
    marcadas = _validar(CON_DETALLE + GENERALES, es_fiscal=es_fiscal)
    assert _codigos(marcadas) == sorted(c for c, _ in CON_DETALLE)


@pytest.mark.parametrize("es_fiscal", [False, True], ids=["PUC Principal", "PUC Fiscal"])
def test_las_cuentas_generales_correctamente_estructuradas_no_se_marcan(es_fiscal):
    assert _validar(GENERALES, es_fiscal=es_fiscal) == []


def test_las_cuentas_maestras_del_grupo_15_nunca_se_marcan_por_sus_nombres_estandar():
    assert _validar([]) == []                                                # los maestros OFICINAS, EDIFICIOS, GARAJES… no generan nada


def test_la_observacion_es_una_recomendacion_con_el_texto_sugerido_la_cuenta_y_el_motivo():
    marcadas = _validar([("15551620", "LOCAL 4 EDIFICIO CENTRAL")])
    assert len(marcadas) == 1
    fila, columna, tipo, mensaje, clave = marcadas[0]
    assert tipo == "advertencia" and columna == "Descripción" and clave == CLAVE        # recomendación: no bloquea
    assert 'La cuenta "15551620" (LOCAL 4 EDIFICIO CENTRAL)' in mensaje                 # código y nombre de la cuenta observada
    assert "unidad específica" in mensaje and "«LOCAL 4»" in mensaje                    # por qué se señala
    assert TEXTO_SUGERIDO in mensaje                                                    # el texto sugerido, completo y sin cambios
    assert fila == len(MAESTROS) + 2                                                    # la fila de Excel de esa cuenta (encabezado = fila 1)


@pytest.mark.parametrize("es_fiscal", [False, True], ids=["PUC Principal", "PUC Fiscal"])
def test_no_se_modifican_ni_se_eliminan_los_datos_del_cliente(es_fiscal):
    _, _, antes, despues = _validar(CON_DETALLE + GENERALES, es_fiscal=es_fiscal, devolver_df=True)
    assert despues.equals(antes) and len(despues) == len(MAESTROS) + len(CON_DETALLE) + len(GENERALES)


def test_una_observacion_por_cuenta_aunque_se_repitan_los_nombres():
    marcadas = _validar([("15551620", "LOCAL 4 EDIFICIO CENTRAL"), ("15551621", "LOCAL 4 EDIFICIO CENTRAL")])
    assert _codigos(marcadas) == ["15551620", "15551621"]


def test_la_regla_no_agrega_ningun_otro_tipo_de_observacion():
    """Un PUC solo con cuentas generales da exactamente lo mismo que antes: ninguna observación de esta regla ni de otra clave nueva."""
    _, errores, _, _ = _validar(GENERALES, devolver_df=True)
    assert not [e for e in errores if "activo" in str(e[4]).lower()]


# ───────────────────────── El Excel marcado ─────────────────────────
def test_el_excel_resalta_la_celda_con_su_comentario_y_no_toca_los_datos():
    import openpyxl
    from openpyxl.styles import PatternFill  # noqa: F401
    cols = ejecutar(["_COLS_PUC"], lambda ns: list(ns["_COLS_PUC"]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIGRACION"
    ws.append(cols)
    filas = [(c, d, "M") for c, d in MAESTROS] + [(c, d, "A") for c, d in CON_DETALLE + GENERALES]
    for c, d, t in filas:
        ws.append([c, d, t] + [None] * (len(cols) - 3))
    datos = io.BytesIO()
    wb.save(datos)
    marcadas = _validar(CON_DETALLE + GENERALES)

    def marcar(ns):
        return ns["marcar_errores_excel"](io.BytesIO(datos.getvalue()), marcadas, None)

    salida = ejecutar(["marcar_errores_excel"], marcar)
    hoja = openpyxl.load_workbook(io.BytesIO(salida))["MIGRACION"]
    col_desc = cols.index("Descripción") + 1
    marcadas_celdas = {}
    for fila in range(2, hoja.max_row + 1):
        celda = hoja.cell(row=fila, column=col_desc)
        if celda.fill and celda.fill.fill_type == "solid":
            marcadas_celdas[str(hoja.cell(row=fila, column=1).value)] = celda
    assert sorted(marcadas_celdas) == sorted(c for c, _ in CON_DETALLE)                # solo las observadas quedan resaltadas
    for celda in marcadas_celdas.values():
        assert celda.fill.start_color.rgb.endswith("FFCCCC")                           # el mismo rojo de alerta de las demás validaciones
        assert celda.comment and "Activos Fijos" in celda.comment.text and "nivel de detalle específico del activo" in celda.comment.text
    for fila, (c, d, t) in enumerate(filas, start=2):                                  # los datos originales quedaron intactos
        assert (str(hoja.cell(row=fila, column=1).value), hoja.cell(row=fila, column=2).value, hoja.cell(row=fila, column=3).value) == (c, d, t)
    resumen = list(openpyxl.load_workbook(io.BytesIO(salida))["Inconsistencias"].iter_rows(values_only=True))[1:]
    assert sorted(re.search(r'"(\d+)"', r[3]).group(1) for r in resumen) == sorted(c for c, _ in CON_DETALLE)


# ───────────────────────── El PDF ─────────────────────────
def _texto_pdf(errores, plantilla="PUC Principal"):
    pypdf = pytest.importorskip("pypdf")
    reemplazo = {"generar_pdf_inconsistencias": [("except Exception:\n        return None", "except Exception:\n        raise")]}
    pdf = ejecutar(["generar_pdf_inconsistencias"], lambda ns: ns["generar_pdf_inconsistencias"](
        "Reporte de Inconsistencias — Plantillas Maestras", [(plantilla, errores)], "Cliente de prueba", "999000001"), reemplazos=reemplazo)
    assert pdf and pdf[:4] == b"%PDF"
    return re.sub(r"\s+", " ", " ".join((p.extract_text() or "") for p in pypdf.PdfReader(io.BytesIO(pdf)).pages))


@pytest.mark.parametrize("plantilla", ["PUC Principal", "PUC Fiscal"])
def test_el_pdf_trae_la_regla_la_descripcion_la_recomendacion_y_cada_cuenta(plantilla):
    texto = _texto_pdf(_validar(CON_DETALLE + GENERALES, es_fiscal=plantilla == "PUC Fiscal"), plantilla)
    assert TITULO in texto                                                              # regla de validación incumplida
    assert "nivel de detalle específico del activo" in texto and "submódulo de Activos Fijos" in texto   # descripción
    assert "Valida si esta clasificación debe manejarse de forma general en el maestro contable" in texto  # recomendación
    for codigo, desc in CON_DETALLE:                                                    # código y nombre de cada cuenta observada
        assert f"{codigo} — {desc[:1].upper() + desc[1:].lower()}" in texto, codigo
    for codigo, _ in GENERALES:                                                         # y ninguna cuenta general
        assert f"{codigo} —" not in texto, codigo


def test_el_excel_y_el_pdf_muestran_las_mismas_cuentas():
    marcadas = _validar(CON_DETALLE + GENERALES)
    en_excel = set(_codigos(marcadas))
    texto = _texto_pdf(marcadas)
    en_pdf = set(re.findall(r"\b(15\d{6,})\s—", texto))
    assert en_excel == en_pdf == {c for c, _ in CON_DETALLE}


# ───────────────────────── En pantalla ─────────────────────────
def _xlsx_puc(auxiliares, fiscal=False):
    import openpyxl
    cols = ejecutar(["_COLS_PUC"], lambda ns: list(ns["_COLS_PUC"])) + (["Equivalencia"] if fiscal else [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIGRACION"
    ws.append(cols)
    for c, d in MAESTROS:
        ws.append([c, d, "M"] + [None] * (len(cols) - 3))
    for c, d in auxiliares:
        ws.append([c, d, "A"] + [None] * (len(cols) - 3))
    datos = io.BytesIO()
    wb.save(datos)
    return datos.getvalue()


def _texto_pantalla(app):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", " ".join(m.value for m in app.markdown)))


def _abrir_etapa1():
    pytest.importorskip("openpyxl")
    from streamlit.testing.v1 import AppTest
    from _cargador import llenar_empresa
    app = AppTest.from_file(os.path.join(RAIZ, "validador_sinco_ayf.py"), default_timeout=180)
    app.run()
    next(b for b in app.button if b.label == "➡️ Iniciar validación").click().run()
    radio = app.radio[0]
    radio.set_value([o for o in radio.options if "Ya tengo mis plantillas" in o][0]).run()
    llenar_empresa(app)
    return app


def _subir(app, base, datos):
    next(f for f in app.file_uploader if f.key == base or f.key.startswith(base + "_")).set_value([("puc.xlsx", datos, XLSX)]).run()


def test_la_pantalla_muestra_la_recomendacion_del_puc_principal_con_detalle_individual():
    app = _abrir_etapa1()
    _subir(app, "puc_principal", _xlsx_puc(CON_DETALLE + GENERALES))
    assert not app.exception, [str(e.value) for e in app.exception]
    assert TITULO in _texto_pantalla(app)


def test_la_pantalla_no_muestra_la_recomendacion_si_el_puc_es_general():
    app = _abrir_etapa1()
    _subir(app, "puc_principal", _xlsx_puc(GENERALES))
    assert not app.exception, [str(e.value) for e in app.exception]
    assert TITULO not in _texto_pantalla(app)


def test_la_pantalla_muestra_la_recomendacion_tambien_en_el_puc_fiscal():
    app = _abrir_etapa1()
    _subir(app, "puc_principal", _xlsx_puc(GENERALES))
    _subir(app, "puc_fiscal", _xlsx_puc(CON_DETALLE, fiscal=True))
    assert not app.exception, [str(e.value) for e in app.exception]
    assert TITULO in _texto_pantalla(app)


# ───────────────────────── La estructura general de referencia (entregada por Andrea) ─────────────────────────
ESTRUCTURA_GENERAL = """15	PROPIEDAD,PLANTA Y EQUIPO
1504	TERRENOS
150401	LOTE
15040101	LOTE PROYECTOS
1508	CONSTRUCCIONES EN CURSO
150801	CONSTRUCCIONES Y EDIFICACIONES EN CURSO 
15080101	CONSTRUCCIONES Y EDIFICACIONES EN CURSO 
1512	MAQUINARIA Y EQUIPOS EN MONTAJE
151205	MAQUINARIA Y EQUIPO
15120501	MAQUINARIA Y EQUIPO
1516	CONSTRUCCIONES Y EDIFICACIONES
151605	EDIFICIOS
15160501	EDIFICIOS 
151610	OFICINAS
15161001	OFICINAS
1520	MAQUINARIA Y EQUIPO
1524	EQUIPO DE OFICINA
152405	MUEBLES Y ENSERES
1528	EQUIPO DE COMPUTACION Y COMUNICACION
1540	FLOTA Y EQUIPO DE TRANSPORTE
154005	AUTOS,CAMIONETAS Y CAMPEROS
15400501	CAMIONETAS
1592	DEPRECIACION ACUMULADA
159205	CONSTRUCCIONES Y EDIFICACIONES
15920501	CONSTRUCCIONES Y EDIFICACIONES
159210	MAQUINARIA Y EQUIPO
15921001	MAQUINARIA Y EQUIPO
159215	EQUIPO DE OFICINA
15921501	EQUIPO DE OFICINA
159220	EQUIPO DE COMPUTACION Y COMUNICACION
15922001	EQUIPO DE COMPUTACION Y COMUNICACION
159235	FLOTA  Y EQUIPO DE TRANSPORTE
15923501	FLOTA  Y EQUIPO DE TRANSPORTE
159240	FLOTA Y EQUIPO  FLUVIAL Y/O MARITIMO
15924001	FLOTA Y EQUIPO  FLUVIAL Y/O MARITIMO
16	INTANGIBLES
1605	CREDITOS MERCANTIL"""


def _validar_estructura(extra=(), es_fiscal=False):
    """Corre validar_puc sobre la estructura de referencia completa (+ cuentas extra) y devuelve [(código, descripción)] de las cuentas marcadas."""
    import pandas as pd

    def correr(ns):
        cols = list(ns["_COLS_PUC"]) + (["Equivalencia"] if es_fiscal and "Equivalencia" not in ns["_COLS_PUC"] else [])
        regs = [{**{c: "" for c in cols}, "Código": c, "Descripción": d, "Tipo": "A" if len(c) >= 8 else "M"}
                for c, d in [l.split("\t", 1) for l in ESTRUCTURA_GENERAL.split("\n")] + list(extra)]
        return ns["validar_puc"](pd.DataFrame(regs), es_fiscal=es_fiscal)
    errores = ejecutar(["validar_puc", "_COLS_PUC"], correr)
    return [re.search(r'La cuenta "(\d+)" \((.*?)\)', e[3]).groups() for e in errores if e[4] == CLAVE]


@pytest.mark.parametrize("es_fiscal", [False, True], ids=["PUC Principal", "PUC Fiscal"])
def test_la_estructura_general_de_referencia_no_genera_ninguna_observacion(es_fiscal):
    """La forma correcta: cada auxiliar lleva el nombre general de su cuenta padre (EDIFICIOS bajo EDIFICIOS, OFICINAS bajo OFICINAS, CAMIONETAS bajo
    AUTOS, CAMIONETAS Y CAMPEROS…). Nada de esto se debe marcar."""
    assert _validar_estructura(es_fiscal=es_fiscal) == []


@pytest.mark.parametrize("codigo,desc", [("15160502", "APARTAMENTOS"), ("15160503", "LOCALES"), ("15160504", "APTO 401 MIRAMAR"),
                                         ("15161002", "OFICINA 502"), ("15040102", "CALLE 128")])
def test_en_esa_misma_estructura_el_detalle_individual_si_se_marca(codigo, desc):
    assert (codigo, desc) in _validar_estructura([(codigo, desc)])


@pytest.mark.parametrize("codigo,desc", [("15400503", "CAMIONETAS"), ("15160505", "EDIFICIOS ADMINISTRATIVOS"), ("15240510", "EQUIPO DE OFICINA")])
def test_en_esa_misma_estructura_un_nombre_general_adicional_no_se_marca(codigo, desc):
    assert _validar_estructura([(codigo, desc)]) == []
