"""Pruebas del archivo «lista para migrar» con SOLO la hoja MIGRACION.

Las plantillas traen, además de MIGRACION (la que se importa a SINCO), hojas de ayuda: Descripción y los catálogos de bancos, ciudades, CIIU… Al llegar a «lista para
migrar» la app entrega el archivo con únicamente la hoja MIGRACION (la plantilla de Terceros pasa de 421 a 11 KB).

Qué se protege:
  * Se quitan todas las hojas menos MIGRACION, y la hoja MIGRACION queda con exactamente los mismos valores, encabezados y dimensiones.
  * Los nombres definidos que apuntaban a hojas quitadas (COD_BANCOS, DV…) y las listas desplegables que los usaban se eliminan para que no queden rotos; los que apuntan a
    MIGRACION (MIGRADOR, Migracion) se conservan, y también las listas con valores propios («S,N»).
  * Las fórmulas se reemplazan por su valor ya calculado: al guardar con openpyxl pierden ese valor y un programa que lea sin recalcular vería celdas vacías.
  * Si el archivo ya trae solo esa hoja, o no tiene hoja MIGRACION, o no se puede leer, se entregan los MISMOS bytes. Si una fórmula lee otra hoja y nunca se calculó, el archivo se
    deja completo antes que dañarlo.
  * La caja verde de «lista para migrar» lo dice, y se combina con el aviso del formato número.

Cómo correrlas (desde la raíz del repositorio):  pip install pytest openpyxl   y luego   pytest pruebas
"""
import io
import os
import re
import zipfile

import pytest

from _cargador import ejecutar

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FUNCIONES = ["solo_hoja_migracion", "_formulas_a_valores", "_preparar_lista_para_migrar", "_normalizar_cacheado", "_solo_hoja_migracion_cacheado", "normalizar_formato_numero",
             "numeros_guardados_como_texto", "errores_numero_como_texto", "_numero_seguro", "_parece_numero", "COLUMNAS_NUMERO", "_sin_tildes", "_aviso_formato_numero",
             "_aviso_hojas_migracion", "_bloque_listas_para_migrar"]
COLS_SI = ["Cuenta Contable", "NIT", "Centro Costos", "Débito", "Crédito"]


def _libro(filas=(), hojas_extra=("Descripción", "BANCOS"), nombres=None, validaciones=None, formulas=None, columnas=COLS_SI):
    """Un libro con la hoja MIGRACION (encabezado + filas) y hojas de ayuda. nombres: {nombre: destino}; validaciones: [(rango, formula1)]; formulas: {celda: fórmula}."""
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIGRACION"
    ws.append(columnas)
    for f in filas:
        ws.append(f)
    for celda, formula in (formulas or {}).items():
        ws[celda] = formula
    for h in hojas_extra:
        wb.create_sheet(h).append(["Codigo", "Entidad"])
    for nombre, destino in (nombres or {}).items():
        wb.defined_names[nombre] = DefinedName(nombre, attr_text=destino)
    for rango, formula in (validaciones or []):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(rango)
        ws.add_data_validation(dv)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _con_valores_calculados(contenido, valores):
    """Escribe en el XML de la hoja el valor calculado de cada fórmula ({celda: valor}), como lo guarda Excel (openpyxl no lo hace)."""
    entrada = zipfile.ZipFile(io.BytesIO(contenido))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for info in entrada.infolist():
            datos = entrada.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = datos.decode("utf-8")
                for celda, valor in valores.items():
                    xml, n = re.subn(rf'(<c r="{celda}"[^>]*>)(<f>[^<]*</f>)(?:<v\s*/>|<v></v>)', rf"\g<1>\g<2><v>{valor}</v>", xml)
                    assert n == 1, celda
                datos = xml.encode("utf-8")
            z.writestr(info, datos)
    return salida.getvalue()


def _quitar(contenido):
    return ejecutar(FUNCIONES, lambda ns: ns["solo_hoja_migracion"](contenido))


def _hojas(contenido):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(contenido)).sheetnames


def _valores(contenido):
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)["MIGRACION"]
    return [[(None if c.value in (None, "") else c.value) for c in fila] for fila in ws.iter_rows()]


# ───────────────────────── Se quitan las hojas de ayuda ─────────────────────────
def test_queda_solo_la_hoja_migracion_y_dice_cuales_quito():
    contenido = _libro([["13050501", "900123456", None, 1000, 0]], hojas_extra=("Descripción", "BANCOS", "CIUDADES-DEPARTAMENTOS"))
    nuevo, quitadas = _quitar(contenido)
    assert _hojas(nuevo) == ["MIGRACION"] and quitadas == ["Descripción", "BANCOS", "CIUDADES-DEPARTAMENTOS"]


def test_la_hoja_migracion_queda_con_los_mismos_valores_encabezados_y_dimensiones():
    import datetime
    filas = [["13050501", "900123456", None, 1000, 0], ["21050501", "900123457", "01", 0, 1000.75], ["13050502", "900123458", None, "texto", datetime.datetime(2026, 1, 15)]]
    contenido = _libro(filas)
    nuevo, _ = _quitar(contenido)
    antes, despues = _valores(contenido), _valores(nuevo)
    assert antes == despues and len(antes) == 4 and despues[0] == COLS_SI
    assert [type(v) for v in despues[2]] == [type(v) for v in antes[2]]                  # los números siguen siendo números, las fechas fechas


def test_si_el_libro_estaba_en_otra_hoja_activa_queda_abierta_la_migracion():
    nuevo, _ = _quitar(_libro())
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(nuevo))
    assert wb.active.title == "MIGRACION"


# ───────────────────────── Nombres definidos y listas desplegables ─────────────────────────
def test_los_nombres_que_apuntaban_a_hojas_quitadas_se_eliminan_y_los_de_migracion_se_conservan():
    import openpyxl
    contenido = _libro(nombres={"COD_BANCOS": "BANCOS!$A:$A", "DV": "'Dígito Verificación'!$A:$E", "CIUDADES": "'BANCOS'!$A:$A", "MIGRADOR": "MIGRACION!$B:$AO", "Migracion": "MIGRACION!$A:$J"},
                       hojas_extra=("Descripción", "BANCOS", "Dígito Verificación"))
    nuevo, _ = _quitar(contenido)
    nombres = {n: d.attr_text for n, d in openpyxl.load_workbook(io.BytesIO(nuevo)).defined_names.items()}
    assert nombres == {"MIGRADOR": "MIGRACION!$B:$AO", "Migracion": "MIGRACION!$A:$J"}


def test_las_listas_desplegables_que_usaban_hojas_o_nombres_quitados_se_eliminan_y_las_propias_se_conservan():
    import openpyxl
    contenido = _libro(nombres={"COD_BANCOS": "BANCOS!$A:$A"}, validaciones=[("D2:D50", "COD_BANCOS"), ("E2:E50", "BANCOS!$A$2:$A$30"), ("C2:C50", '"S,N"'), ("B2:B50", "MIGRACION!$A$2:$A$9")])
    nuevo, _ = _quitar(contenido)
    ws = openpyxl.load_workbook(io.BytesIO(nuevo))["MIGRACION"]
    assert sorted(dv.formula1 for dv in ws.data_validations.dataValidation) == ['"S,N"', "MIGRACION!$A$2:$A$9"]


# ───────────────────────── Cuándo se entrega el archivo tal cual ─────────────────────────
def test_si_ya_trae_solo_la_hoja_migracion_se_entregan_los_mismos_bytes():
    contenido = _libro(hojas_extra=())
    nuevo, quitadas = _quitar(contenido)
    assert nuevo == contenido and quitadas == []


def test_si_no_tiene_una_hoja_llamada_migracion_no_se_toca():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Datos"
    wb.create_sheet("Otra")
    b = io.BytesIO()
    wb.save(b)
    nuevo, quitadas = _quitar(b.getvalue())
    assert nuevo == b.getvalue() and quitadas == []


def test_un_archivo_que_no_se_puede_leer_se_devuelve_igual():
    nuevo, quitadas = _quitar(b"esto no es un excel")
    assert nuevo == b"esto no es un excel" and quitadas == []


def test_aplicarlo_dos_veces_no_cambia_nada_la_segunda():
    primero, _ = _quitar(_libro([["13050501", "900123456", None, 1000, 0]]))
    segundo, quitadas = _quitar(primero)
    assert segundo == primero and quitadas == []


# ───────────────────────── Fórmulas ─────────────────────────
def test_las_formulas_se_reemplazan_por_su_valor_calculado():
    """Caso real: el ejemplo de Movimientos trae una fila de totales con =SUM(...). Sin esto, al guardar queda sin valor."""
    contenido = _con_valores_calculados(_libro([["13050501", "900123456", None, 1000, 0], ["21050501", "900123457", None, 390000, 0]], formulas={"D4": "=SUM(D2:D3)", "E4": "=SUM(E2:E3)"}),
                                        {"D4": 391000, "E4": 0})
    assert _valores(contenido)[3][3:] == [391000, 0]                                     # el archivo de entrada sí trae los valores calculados
    nuevo, quitadas = _quitar(contenido)
    assert quitadas and _valores(nuevo)[3][3:] == [391000, 0]
    import openpyxl
    assert not [c for f in openpyxl.load_workbook(io.BytesIO(nuevo))["MIGRACION"].iter_rows() for c in f if isinstance(c.value, str) and c.value.startswith("=")]


def test_una_formula_que_lee_una_hoja_quitada_y_nunca_se_calculo_deja_el_archivo_completo():
    """Sin valor calculado no se puede reemplazar sin dañar: se entrega el archivo con todas sus hojas."""
    contenido = _libro([["13050501", "900123456", None, 1000, 0]], formulas={"C2": "=BANCOS!A2"})
    nuevo, quitadas = _quitar(contenido)
    assert nuevo == contenido and quitadas == []


def test_una_formula_que_lee_una_hoja_quitada_y_si_estaba_calculada_se_reemplaza_por_su_valor():
    contenido = _con_valores_calculados(_libro([["13050501", "900123456", None, 1000, 0]], formulas={"C2": "=BANCOS!A2"}), {"C2": 51})
    nuevo, quitadas = _quitar(contenido)
    assert quitadas and _hojas(nuevo) == ["MIGRACION"] and _valores(nuevo)[1][2] == 51


# ───────────────────────── Junto con el formato número ─────────────────────────
def _preparar(contenido, tipo):
    return ejecutar(FUNCIONES, lambda ns: ns["_preparar_lista_para_migrar"](contenido, tipo))


def test_la_lista_para_migrar_aplica_el_formato_numero_y_deja_solo_la_hoja_migracion():
    contenido = _libro([["13050501", "900123456", None, "1500000", "0"], ["13050502", "900123457", None, 0, 1500000]])
    datos, conteo, hojas = _preparar(contenido, "si")
    assert conteo == {"Débito": 1, "Crédito": 1} and hojas == ["Descripción", "BANCOS"] and _hojas(datos) == ["MIGRACION"]
    assert _valores(datos)[1][3:] == [1500000, 0] and all(type(v) is int for f in _valores(datos)[1:] for v in f[3:])


def test_sin_tipo_solo_se_quitan_las_hojas_como_en_la_etapa_1():
    contenido = _libro([["13050501", "900123456", None, "1500000", "0"]])
    datos, conteo, hojas = _preparar(contenido, None)
    assert conteo == {} and hojas == ["Descripción", "BANCOS"] and _valores(datos)[1][3:] == ["1500000", "0"]      # el texto no se toca fuera de las Etapas 2 y 3


def test_un_archivo_ya_correcto_y_con_una_sola_hoja_se_entrega_con_los_mismos_bytes():
    contenido = _libro([["13050501", "900123456", None, 1000, 0]], hojas_extra=())
    datos, conteo, hojas = _preparar(contenido, "si")
    assert datos == contenido and conteo == {} and hojas == []


# ───────────────────────── El mensaje ─────────────────────────
def _aviso(quitadas):
    return ejecutar(FUNCIONES, lambda ns: ns["_aviso_hojas_migracion"](quitadas))


def test_sin_hojas_quitadas_no_hay_aviso():
    assert _aviso([]) == "" and _aviso([("Terceros", [])]) == ""


def test_el_aviso_dice_que_solo_queda_la_hoja_migracion_y_cuales_se_quitaron():
    texto = re.sub(r"<[^>]+>", "", _aviso([("Terceros", ["Descripción", "BANCOS"]), ("PUC Principal", []), ("Centros de Costos", ["Descripción"])]))
    assert "solo la hoja MIGRACION" in texto and "la que se importa" in texto
    assert "Terceros: Descripción, BANCOS" in texto and "Centros de Costos: Descripción" in texto and "PUC Principal" not in texto


def _bloque(listas):
    from streamlit.testing.v1 import AppTest

    def script(bloque, datos):
        bloque(datos, "e1")

    def correr(ns):
        import datetime as _dt
        import streamlit as _st
        ns["st"], ns["re"], ns["datetime"] = _st, re, _dt.datetime
        at = AppTest.from_function(script, args=(ns["_bloque_listas_para_migrar"], listas), default_timeout=60).run()
        assert not at.exception, [str(e.value)[:200] for e in at.exception]
        texto = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", " ".join(m.value for m in at.markdown if "<style>" not in m.value)))
        return texto.replace(" :", ":")
    return ejecutar(FUNCIONES, correr)


def test_la_caja_verde_dice_que_dejo_solo_la_hoja_migracion():
    texto = _bloque([("Terceros", "Terceros", _libro([["13050501", "900123456", None, 1000, 0]], hojas_extra=("Descripción", "BANCOS")))])
    assert "Plantilla lista para migrar" in texto and "el archivo que validaste, listo para migrar" in texto
    assert "solo la hoja MIGRACION" in texto and "Terceros: Descripción, BANCOS" in texto and "sin ningún cambio" not in texto


def test_la_caja_verde_combina_las_hojas_quitadas_con_el_formato_numero():
    texto = _bloque([("Saldos Iniciales Principal", "SI Principal", _libro([["13050501", "900123456", None, "1500000", 0]]), "si")])
    assert "solo la hoja MIGRACION" in texto and "Le aplicamos el formato número" in texto and "Saldos Iniciales Principal: 1 valores (Débito 1)" in texto


def test_la_caja_verde_sigue_diciendo_el_mismo_archivo_cuando_no_hay_nada_que_cambiar():
    texto = _bloque([("PUC Principal", "PUC Principal", _libro([["1105", "Caja"]], hojas_extra=(), columnas=["Código", "Descripción"]))])
    assert "el mismo archivo que validaste, sin ningún cambio" in texto and "solo la hoja" not in texto


# ───────────────────────── Con las plantillas reales del repositorio ─────────────────────────
@pytest.mark.parametrize("nombre,hojas_quitadas", [("Plantilla_Terceros.xlsx", 6), ("Ejemplo_Terceros.xlsx", 6), ("Plantilla_Centros_de_Costos.xlsx", 3), ("Plantilla_PUC_Principal_NIF.xlsx", 1),
                                                    ("Ejemplo_Movimientos.xlsx", 1), ("Plantilla_Saldos_CxP.xlsx", 1)])
def test_las_plantillas_oficiales_quedan_con_solo_migracion_y_con_los_mismos_valores(nombre, hojas_quitadas):
    contenido = open(os.path.join(RAIZ, "plantillas", nombre), "rb").read()
    nuevo, quitadas = _quitar(contenido)
    assert len(quitadas) == hojas_quitadas and _hojas(nuevo) == ["MIGRACION"] and _valores(nuevo) == _valores(contenido)
    import openpyxl
    assert all(d.attr_text.startswith("MIGRACION!") for d in openpyxl.load_workbook(io.BytesIO(nuevo)).defined_names.values())     # solo quedan nombres que apuntan a MIGRACION


def test_la_plantilla_de_terceros_baja_de_peso():
    contenido = open(os.path.join(RAIZ, "plantillas", "Plantilla_Terceros.xlsx"), "rb").read()
    assert len(_quitar(contenido)[0]) < len(contenido) / 5


# ───────────────────────── Dónde se aplica ─────────────────────────
def test_todas_las_etapas_pasan_por_el_mismo_bloque_que_deja_solo_la_hoja_migracion():
    src = open(os.path.join(RAIZ, "validador_sinco_ayf.py"), encoding="utf-8").read()
    assert "_datos, _conteo, _hojas = _preparar_lista_para_migrar(_item[2], _item[3] if len(_item) > 3 else None)" in src
    assert src.count("_bloque_listas_para_migrar(") >= 4                                  # la definición y las llamadas de las Etapas 1, 2 y 3
